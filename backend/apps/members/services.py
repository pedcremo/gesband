import csv
import hashlib
import io
import logging
import re
import unicodedata
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from pathlib import Path
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.associations.models import Association

from .catalog import catalog_key, section_codes_by_key
from .catalog import instruments_by_key as catalog_instruments_by_key
from .catalog import seed_instruments
from .models import ImportBatch, Instrument, Member, Section

logger = logging.getLogger(__name__)

MODEL_COLUMNS = {"external_id", "first_name", "last_name", "national_id", "address", "phone", "email", "kind"}
SURNAME_COLUMNS = {"first_surname", "second_surname"}
# Columnas opcionales que se resuelven contra el catálogo de la asociación.
CATALOG_COLUMNS = {"section", "instrument"}
ALLOWED_COLUMNS = MODEL_COLUMNS | SURNAME_COLUMNS | CATALOG_COLUMNS
MAX_IMPORT_ROWS = 1000
MAX_IMPORT_BYTES = 5 * 1024 * 1024

HEADER_ALIASES = {
    "external_id": {"external id", "id externo", "identificador externo", "codigo", "codi", "numero socio"},
    "first_name": {"first name", "nombre", "nom"},
    "last_name": {"last name", "apellidos", "cognoms"},
    "first_surname": {"first surname", "surname 1", "apellido 1", "primer apellido", "cognom 1", "primer cognom"},
    "second_surname": {"second surname", "surname 2", "apellido 2", "segundo apellido", "cognom 2", "segon cognom"},
    "national_id": {"national id", "dni", "nie", "dni nie", "nif"},
    "address": {"address", "direccion", "adreca"},
    "phone": {"phone", "telefono", "telefon", "movil", "mobil"},
    "email": {"email", "correo", "correo electronico", "correu", "correu electronic"},
    "kind": {"kind", "tipo", "tipo de miembro", "tipus", "tipus de membre"},
    "section": {"section", "cuerda", "corda", "seccion", "seccio"},
    "instrument": {"instrument", "instrumento", "instrumentos", "instruments"},
}
# Una celda puede traer varios instrumentos: «Clarinete; Requinto» o «Clarinete / Requinto».
INSTRUMENT_SEPARATORS = re.compile(r"[;/]")
# Un nombre de catálogo que empieza así se interpretaría como fórmula al exportarlo.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
CATALOG_NAME_MAX_LENGTH = 80


def _normalize_header(value):
    ascii_value = "".join(
        character
        for character in unicodedata.normalize("NFKD", str(value))
        if not unicodedata.combining(character)
    )
    return re.sub(r"[^a-z0-9]+", " ", ascii_value.casefold()).strip()


def suggest_import_mapping(headers):
    """Return a conservative mapping for known template and human-readable headers."""
    normalized_aliases = {
        alias: target
        for target, aliases in HEADER_ALIASES.items()
        for alias in aliases | {_normalize_header(target)}
    }
    suggestions = {}
    claimed_targets = set()
    for header in headers:
        target = normalized_aliases.get(_normalize_header(header))
        if target and target not in claimed_targets:
            suggestions[header] = target
            claimed_targets.add(target)
    if "first_surname" in claimed_targets and "last_name" in claimed_targets:
        suggestions = {header: target for header, target in suggestions.items() if target != "last_name"}
    return suggestions


def inspect_import(upload):
    content = upload.read()
    if len(content) > MAX_IMPORT_BYTES:
        raise ValidationError(_("El fichero supera 5 MB."))
    extension = Path(upload.name).suffix.lower()
    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValidationError(_("El CSV debe estar codificado en UTF-8.")) from exc
        reader = csv.reader(io.StringIO(text))
        headers = [str(value or "").strip() for value in next(reader, [])]
        rows = [dict(zip(headers, values)) for values in reader]
    elif extension == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator, [])]
            rows = [dict(zip(headers, values)) for values in iterator]
            workbook.close()
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise ValidationError(_("El fichero XLSX no es válido.")) from exc
    else:
        raise ValidationError(_("Solo se admiten ficheros CSV o XLSX."))
    if not headers or any(not header for header in headers):
        raise ValidationError(_("Todas las columnas deben tener un nombre."))
    if len(set(headers)) != len(headers):
        raise ValidationError(_("Los nombres de columna no pueden repetirse."))
    if not rows:
        raise ValidationError(_("El fichero no contiene filas de datos."))
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValidationError(_("El fichero supera el límite de %(limit)s filas.") % {"limit": MAX_IMPORT_ROWS})
    return content, hashlib.sha256(content).hexdigest(), headers, rows


def parse_import(upload, mapping, association=None):
    """Valida el fichero con su mapeo sin escribir en el censo ni en el catálogo.

    Con `association`, las filas incluyen además los errores y avisos de
    cuerdas e instrumentos frente al catálogo de esa asociación.
    """
    content, content_hash, headers, rows = inspect_import(upload)

    invalid_targets = set(mapping.values()) - ALLOWED_COLUMNS
    if invalid_targets:
        raise ValidationError(
            _("Campos de destino no permitidos: %(fields)s.") % {"fields": ", ".join(sorted(invalid_targets))}
        )
    unknown_sources = set(mapping) - set(headers)
    if unknown_sources:
        raise ValidationError(
            _("Columnas de origen desconocidas: %(fields)s.") % {"fields": ", ".join(sorted(unknown_sources))}
        )
    targets = list(mapping.values())
    if len(targets) != len(set(targets)):
        raise ValidationError(_("Cada campo de destino solo puede asignarse una vez."))
    if "first_name" not in targets:
        raise ValidationError(_("Asigna una columna al campo Nombre."))
    uses_combined_surnames = "last_name" in targets
    uses_split_surnames = bool(SURNAME_COLUMNS.intersection(targets))
    if not uses_combined_surnames and "first_surname" not in targets:
        raise ValidationError(_("Asigna Apellidos (juntos) o Primer apellido."))
    if uses_combined_surnames and uses_split_surnames:
        raise ValidationError(_("Usa Apellidos (juntos) o los apellidos separados, pero no ambos."))
    if "second_surname" in targets and "first_surname" not in targets:
        raise ValidationError(_("El Segundo apellido requiere asignar también el Primer apellido."))
    normalized = []
    external_ids = {}
    for row_number, source in enumerate(rows, start=2):
        data = {target: str(source.get(origin) or "").strip() for origin, target in mapping.items()}
        errors = []
        if not data.get("first_name"):
            errors.append(_("Falta el nombre."))
        if uses_split_surnames:
            first_surname = data.pop("first_surname", "")
            second_surname = data.pop("second_surname", "")
            data["last_name"] = " ".join(part for part in (first_surname, second_surname) if part)
            if not first_surname:
                errors.append(_("Falta el primer apellido."))
        elif not data.get("last_name"):
            errors.append(_("Faltan los apellidos."))
        if data.get("kind") and data["kind"] not in Member.Kind.values:
            errors.append(_("Tipo de miembro no válido."))
        external_id = data.get("external_id")
        if external_id:
            if external_id in external_ids:
                errors.append(_("Identificador externo duplicado en el fichero."))
            else:
                external_ids[external_id] = row_number
        warnings = _parse_catalog_cells(data, errors)
        normalized.append({"row": row_number, "data": data, "errors": errors, "warnings": warnings})
    if association is not None and uses_catalog_columns(mapping):
        normalized, _plan = review_import(association, normalized)
    return content, content_hash, normalized


# --- Cuerdas e instrumentos en la importación --------------------------------
#
# Las cuerdas son comunes y fijas (decisión 0012): se reconocen por su nombre
# en castellano, valenciano o inglés y por los alias de `catalog.py`, y una
# cuerda desconocida invalida la fila. La previsualización solo lee el catálogo
# de instrumentos: informa de los que faltan, pero no los crea. Se crean al
# confirmar, en la misma transacción que las fichas. Siempre se busca en los
# instrumentos de la asociación del lote; los de otra nunca se reutilizan.


def uses_catalog_columns(mapping):
    return bool(CATALOG_COLUMNS.intersection((mapping or {}).values()))


def _clean_catalog_name(value):
    return " ".join(str(value or "").split())


def _catalog_name_errors(name, kind_label):
    errors = []
    if name.startswith(FORMULA_PREFIXES):
        errors.append(
            _("%(kind)s «%(name)s» no puede empezar por =, +, - o @.") % {"kind": kind_label, "name": name}
        )
    if len(name) > CATALOG_NAME_MAX_LENGTH:
        errors.append(
            _("%(kind)s «%(name)s» supera %(limit)s caracteres.")
            % {"kind": kind_label, "name": name[:20] + "…", "limit": CATALOG_NAME_MAX_LENGTH}
        )
    return errors


def _parse_catalog_cells(data, errors):
    """Normaliza las celdas de cuerda e instrumentos de una fila; devuelve avisos.

    Deja `data["section"]` como texto y `data["instruments"]` como lista sin
    repetidos (el primero es el principal). Los errores se añaden a `errors`.
    """
    warnings = []
    if "section" in data:
        data["section"] = _clean_catalog_name(data["section"])
        if data["section"]:
            errors.extend(_catalog_name_errors(data["section"], _("La cuerda")))
    if "instrument" in data:
        names = []
        seen = set()
        for part in INSTRUMENT_SEPARATORS.split(data.pop("instrument")):
            name = _clean_catalog_name(part)
            key = catalog_key(name)
            if not key or key in seen:
                continue
            seen.add(key)
            names.append(name)
            errors.extend(_catalog_name_errors(name, _("El instrumento")))
        data["instruments"] = names
    if data.get("section") and not data.get("instruments"):
        warnings.append(
            _("La cuerda «%(section)s» se ignora: la cuerda de una ficha es la de su instrumento.")
            % {"section": data["section"]}
        )
    return warnings


@dataclass
class ImportCatalogPlan:
    """Qué hay ya en el catálogo de la asociación y qué habría que crear."""

    sections: dict = dataclass_field(default_factory=dict)  # clave -> Section común
    instruments: dict = dataclass_field(default_factory=dict)  # clave -> Instrument existente
    new_instruments: dict = dataclass_field(default_factory=dict)  # clave del plan -> {"name", "entry", "section"}
    name_keys: dict = dataclass_field(default_factory=dict)  # clave de un nombre del fichero -> clave del plan
    row_errors: dict = dataclass_field(default_factory=dict)  # fila -> [mensajes]
    row_warnings: dict = dataclass_field(default_factory=dict)  # fila -> [mensajes]

    def new_instrument_labels(self):
        labels = []
        for item in self.new_instruments.values():
            section = item["section"]
            labels.append(f"{item['name']} ({section.label})" if section else item["name"])
        return labels


def section_lookup():
    """Clave de cualquier nombre o alias de una cuerda común -> Section."""
    sections = {section.code: section for section in Section.objects.all()}
    lookup = {}
    for section in sections.values():
        for name in (section.code, section.name, section.name_ca, section.name_en):
            if name:
                lookup.setdefault(catalog_key(name), section)
    for key, code in section_codes_by_key().items():
        if code in sections:
            lookup.setdefault(key, sections[code])
    return lookup


def instrument_lookup(association):
    """Clave de cualquier nombre o alias -> instrumento de la asociación."""
    instruments = list(Instrument.objects.filter(association=association).select_related("section").order_by("pk"))
    lookup = {}
    for instrument in instruments:
        for name in (instrument.name, instrument.name_ca, instrument.name_en):
            if name:
                lookup.setdefault(catalog_key(name), instrument)
    by_code = {instrument.code: instrument for instrument in instruments if instrument.code}
    for key, entry in catalog_instruments_by_key().items():
        if entry[0] in by_code:
            lookup.setdefault(key, by_code[entry[0]])
    return lookup


def _section_mismatch(name, current, section_label):
    return _(
        "El instrumento «%(instrument)s» ya está en la cuerda «%(current)s» del catálogo, "
        "no en «%(section)s». Corrige el fichero o el catálogo."
    ) % {"instrument": name, "current": current.label, "section": section_label}


def plan_import_catalog(association, rows):
    """Resuelve cuerdas e instrumentos de las filas válidas sin escribir nada.

    Una cuerda que no es una de las comunes, un instrumento que ya existe (o
    está predefinido) bajo otra cuerda, o que el propio fichero pone en dos
    cuerdas distintas, invalida la fila: no se reasigna en silencio.
    """
    plan = ImportCatalogPlan(sections=section_lookup(), instruments=instrument_lookup(association))
    catalog_entries = catalog_instruments_by_key()
    sections_by_code = {section.code: section for section in plan.sections.values()}
    valid_sections = ", ".join(section.label for section in sorted(sections_by_code.values(), key=lambda s: s.order))

    for item in rows:
        names = item["data"].get("instruments") or []
        if item["errors"] or not names:
            continue
        section_label = item["data"].get("section", "")
        section = plan.sections.get(catalog_key(section_label)) if section_label else None
        if section_label and section is None:
            plan.row_errors[item["row"]] = [
                _("La cuerda «%(section)s» no existe. Usa una de estas: %(sections)s.")
                % {"section": section_label, "sections": valid_sections}
            ]
            continue
        errors = []
        warnings = []
        planned = {}
        for name in names:
            key = catalog_key(name)
            existing = plan.instruments.get(key)
            if existing is not None:
                if section and existing.section and existing.section != section:
                    errors.append(_section_mismatch(existing.name, existing.section, section_label))
                elif section and not existing.section:
                    warnings.append(
                        _("El instrumento «%(instrument)s» no tiene cuerda en el catálogo; no se le asigna «%(section)s».")
                        % {"instrument": existing.name, "section": section_label}
                    )
                continue
            entry = catalog_entries.get(key)
            predefined_section = sections_by_code.get(entry[1]) if entry else None
            if section and predefined_section and predefined_section != section:
                errors.append(_section_mismatch(entry[2], predefined_section, section_label))
                continue
            plan_key = f"catalog:{entry[0]}" if entry else key
            pending = plan.new_instruments.get(plan_key)
            if pending and section and pending["section"] and pending["section"] != section:
                errors.append(
                    _("El instrumento «%(instrument)s» aparece en el fichero con las cuerdas «%(first)s» y «%(section)s».")
                    % {"instrument": pending["name"], "first": pending["section"].label, "section": section_label}
                )
                continue
            planned[key] = (plan_key, entry[2] if entry else name, entry, predefined_section)
        if errors:
            plan.row_errors[item["row"]] = errors
            continue
        if warnings:
            plan.row_warnings[item["row"]] = warnings
        for key, (plan_key, name, entry, predefined_section) in planned.items():
            pending = plan.new_instruments.setdefault(
                plan_key, {"name": name, "entry": entry, "section": predefined_section}
            )
            if section and not pending["section"]:
                pending["section"] = section
            plan.name_keys[key] = plan_key
    return plan


def seed_default_instruments(association):
    """Añade a la asociación los instrumentos predefinidos que le falten."""
    with transaction.atomic():
        Association.objects.select_for_update().get(pk=association.pk)
        return seed_instruments(association, Instrument, Section)


def _merge_messages(current, extra):
    return list(current) + [message for message in extra if message not in current]


def review_import(association, rows):
    """Filas con los errores y avisos del catálogo añadidos, y el plan usado."""
    plan = plan_import_catalog(association, rows)
    reviewed = []
    for item in rows:
        row = item["row"]
        reviewed.append(
            {
                **item,
                "errors": _merge_messages(item["errors"], plan.row_errors.get(row, [])),
                "warnings": _merge_messages(item.get("warnings", []), plan.row_warnings.get(row, [])),
            }
        )
    return reviewed, plan


def _create_planned_catalog(association, plan):
    """Crea los instrumentos que faltan y devuelve clave -> Instrument."""
    instruments = dict(plan.instruments)
    for plan_key, item in plan.new_instruments.items():
        entry = item["entry"]
        predefined = {"code": entry[0], "name_ca": entry[3], "name_en": entry[4]} if entry else {}
        instruments[plan_key] = Instrument.objects.create(
            association=association, name=item["name"], section=item["section"], **predefined
        )
    for key, plan_key in plan.name_keys.items():
        instruments[key] = instruments[plan_key]
    return instruments


@transaction.atomic
def confirm_import(batch):
    locked = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if locked.status == ImportBatch.Status.CONFIRMED:
        return locked.result
    if locked.status != ImportBatch.Status.PREVIEW:
        raise ValidationError(_("Genera la previsualización antes de confirmar."))
    rows = locked.preview
    uses_catalog = uses_catalog_columns(locked.mapping)
    if uses_catalog:
        # Bloquear la asociación serializa la creación de catálogo entre lotes
        # concurrentes; el plan se recalcula porque el catálogo pudo cambiar.
        association = Association.objects.select_for_update().get(pk=locked.association_id)
        rows, plan = review_import(association, rows)
    if not any(not item["errors"] for item in rows):
        raise ValidationError(_("No hay filas válidas que confirmar. Corrige el mapeo o el fichero."))
    instruments_by_key = _create_planned_catalog(association, plan) if uses_catalog else {}
    created = 0
    updated = 0
    rejected = 0
    report = []
    for item in rows:
        report_item = {**item}
        if item["errors"]:
            rejected += 1
            report_item["action"] = "rejected"
            report.append(report_item)
            continue
        data = item["data"].copy()
        external_id = data.pop("external_id", "")
        data.pop("section", None)
        instrument_names = data.pop("instruments", None) or []
        defaults = {**data, "kind": data.get("kind") or Member.Kind.MUSICIAN}
        if external_id:
            member, was_created = Member.objects.update_or_create(
                association=locked.association, external_id=external_id, defaults=defaults
            )
        else:
            member = Member.objects.create(association=locked.association, **defaults)
            was_created = True
        if instrument_names:
            # La celda manda: al reimportar se sustituyen los instrumentos de la
            # ficha sin duplicarlos; una celda vacía no toca los que ya tenga.
            instruments = [instruments_by_key[catalog_key(name)] for name in instrument_names]
            _sync_instruments(member, instruments, instruments[0])
        created += int(was_created)
        updated += int(not was_created)
        report_item["action"] = "created" if was_created else "updated"
        report.append(report_item)
    locked.status = ImportBatch.Status.CONFIRMED
    locked.confirmed_at = timezone.now()
    locked.preview = report
    locked.result = {"created": created, "updated": updated, "rejected": rejected}
    if uses_catalog:
        locked.result["instruments_created"] = plan.new_instrument_labels()
    source_name = locked.source_file.name
    source_storage = locked.source_file.storage
    locked.source_file = ""
    locked.save(update_fields=["status", "confirmed_at", "preview", "result", "source_file"])
    if source_name:
        transaction.on_commit(lambda: source_storage.delete(source_name))
    return locked.result


# --- Mantenimiento de fichas desde el panel y la API (decision 0009) ---------
#
# Las reglas viven aqui y no en las vistas: el panel y la API llaman a estas
# funciones, que comprueban permiso, asociacion y coherencia de las entidades
# relacionadas antes de escribir nada.

MEMBER_EDITABLE_FIELDS = (
    "external_id",
    "first_name",
    "last_name",
    "national_id",
    "address",
    "phone",
    "email",
    "kind",
)


def _ensure_manager(actor, association):
    from apps.associations.permissions import ensure_permission

    ensure_permission(actor, association, "members.manage")


def _clean_instruments(association, instruments, primary_instrument):
    """Instrumentos de la propia asociacion; uno ajeno invalida la peticion."""
    from .models import Instrument

    instruments = list(instruments or [])
    if any(item.association_id != association.id for item in instruments):
        raise ValidationError({"instruments": _("El instrumento pertenece a otra asociación.")})
    if primary_instrument is not None:
        if primary_instrument.association_id != association.id:
            raise ValidationError({"primary_instrument": _("El instrumento pertenece a otra asociación.")})
        if primary_instrument not in instruments:
            instruments.append(primary_instrument)
    # Se vuelven a leer de la base para no fiarse de objetos construidos fuera.
    ids = {item.pk for item in instruments}
    if Instrument.objects.filter(association=association, pk__in=ids).count() != len(ids):
        raise ValidationError({"instruments": _("El instrumento pertenece a otra asociación.")})
    return instruments


def _clean_positions(positions):
    from .models import MemberPosition

    positions = set(positions or [])
    unknown = positions - set(MemberPosition.Position.values)
    if unknown:
        raise ValidationError({"positions": _("Cargo no válido.")})
    return positions


def _check_external_id(association, external_id, exclude_pk=None):
    if not external_id:
        return
    duplicates = Member.objects.filter(association=association, external_id=external_id)
    if exclude_pk:
        duplicates = duplicates.exclude(pk=exclude_pk)
    if duplicates.exists():
        raise ValidationError({"external_id": _("Ya hay otra ficha con este identificador externo en la asociación.")})


def _sync_instruments(member, instruments, primary_instrument):
    from .models import MemberInstrument

    wanted = {item.pk for item in instruments}
    member.member_instruments.exclude(instrument_id__in=wanted).delete()
    # Primero se quita la marca de principal para no chocar con la restriccion
    # de un solo instrumento principal por ficha.
    member.member_instruments.update(is_primary=False)
    existing = set(member.member_instruments.values_list("instrument_id", flat=True))
    MemberInstrument.objects.bulk_create(
        [MemberInstrument(member=member, instrument_id=pk) for pk in wanted - existing]
    )
    if primary_instrument is not None:
        member.member_instruments.filter(instrument=primary_instrument).update(is_primary=True)


def _sync_positions(member, positions):
    from .models import MemberPosition

    member.positions.exclude(position__in=positions).delete()
    existing = set(member.positions.values_list("position", flat=True))
    MemberPosition.objects.bulk_create(
        [MemberPosition(member=member, position=position) for position in positions - existing]
    )


@transaction.atomic
def create_member(*, association, actor, data, instruments=None, primary_instrument=None, positions=None):
    """Alta de una ficha. Nace activa salvo que se indique otro estado."""
    _ensure_manager(actor, association)
    instruments = _clean_instruments(association, instruments, primary_instrument)
    positions = _clean_positions(positions)
    values = {field: data[field] for field in MEMBER_EDITABLE_FIELDS if field in data}
    _check_external_id(association, values.get("external_id"))
    status = data.get("status") or Member.Status.ACTIVE
    member = Member(association=association, created_by=actor, updated_by=actor, status=status, **values)
    if status != Member.Status.ACTIVE:
        member.status_changed_at = timezone.now()
        member.status_changed_by = actor
    member.full_clean(exclude=["photo"])
    member.save()
    _sync_instruments(member, instruments, primary_instrument)
    _sync_positions(member, positions)
    return member


@transaction.atomic
def update_member(*, member, actor, data, instruments=None, primary_instrument=None, positions=None):
    """Edita los datos de la ficha y, si se indican, instrumentos y cargos.

    `instruments` y `positions` en `None` significan «no tocar»; una lista vacía
    los deja vacíos. El estado pasa por `set_member_status` para dejar rastro.
    """
    locked = Member.objects.select_for_update().select_related("association").get(pk=member.pk)
    association = locked.association
    _ensure_manager(actor, association)
    values = {field: data[field] for field in MEMBER_EDITABLE_FIELDS if field in data}
    _check_external_id(association, values.get("external_id"), exclude_pk=locked.pk)
    if instruments is not None or primary_instrument is not None:
        instruments = _clean_instruments(association, instruments, primary_instrument)
    if positions is not None:
        positions = _clean_positions(positions)
    for field, value in values.items():
        setattr(locked, field, value)
    locked.updated_by = actor
    locked.full_clean(exclude=["photo"])
    locked.save()
    if instruments is not None:
        _sync_instruments(locked, instruments, primary_instrument)
    if positions is not None:
        _sync_positions(locked, positions)
    if "status" in data and data["status"] and data["status"] != locked.status:
        locked = set_member_status(member=locked, actor=actor, status=data["status"])
    return locked


@transaction.atomic
def set_member_status(*, member, actor, status):
    """Baja (inactivo) o reactivación. No borra convocatorias ni asistencias.

    Una ficha inactiva deja de aparecer al convocar y al invitar, pero su
    historial sigue consultable. La cuenta vinculada, si la hay, no se toca:
    su acceso se gestiona aparte (`apps.accounts.services`).
    """
    if status not in Member.Status.values:
        raise ValidationError({"status": _("Estado no válido.")})
    locked = Member.objects.select_for_update().select_related("association").get(pk=member.pk)
    _ensure_manager(actor, locked.association)
    if locked.status != status:
        locked.status = status
        locked.status_changed_at = timezone.now()
        locked.status_changed_by = actor
        locked.updated_by = actor
        locked.save(update_fields=["status", "status_changed_at", "status_changed_by", "updated_by", "updated_at"])
    return locked


def member_history(member):
    """Lo que impide borrar una ficha sin perder historial, con su recuento."""
    from apps.activities.models import Invitation
    from apps.polls.models import PollRecipient
    from apps.transport.models import Transport, TransportAssignment

    return {
        "invitations": Invitation.objects.filter(member=member).count(),
        "poll_recipients": PollRecipient.objects.filter(member=member).count(),
        "transport_assignments": TransportAssignment.objects.filter(member=member).count(),
        "driven_transports": Transport.objects.filter(driver=member).count(),
    }


def member_deletion_blockers(member):
    """Motivos por los que la ficha no se puede borrar; vacío si se puede.

    Regla de la decisión 0009: solo se borra una ficha que no ha dejado rastro
    (sin convocatorias ni asistencia, sin encuestas, sin transporte) y que no
    tiene una cuenta vinculada. En los demás casos, la baja es la vía.
    """
    blockers = []
    history = member_history(member)
    if history["invitations"]:
        blockers.append(_("Tiene convocatorias o asistencia registradas."))
    if history["poll_recipients"]:
        blockers.append(_("Ha sido destinataria de encuestas."))
    if history["transport_assignments"] or history["driven_transports"]:
        blockers.append(_("Tiene transporte asignado."))
    if member.account_id:
        blockers.append(_("Tiene una cuenta de acceso vinculada."))
    return blockers


@transaction.atomic
def delete_member(*, member, actor):
    """Borrado físico: solo `admin` y solo si no rompe ningún historial."""
    from apps.associations.permissions import ensure_permission

    locked = Member.objects.select_for_update().select_related("association").get(pk=member.pk)
    ensure_permission(actor, locked.association, "members.delete")
    blockers = member_deletion_blockers(locked)
    if blockers:
        raise ValidationError(
            [_("No se puede borrar esta ficha sin perder su historial. Dala de baja en su lugar.")] + blockers
        )
    photo_name = locked.photo.name
    photo_storage = locked.photo.storage
    member_id = locked.pk
    # Las invitaciones de acceso no aceptadas caen con la ficha (CASCADE).
    locked.delete()
    if photo_name:
        transaction.on_commit(lambda: photo_storage.delete(photo_name))
    logger.info("Ficha %s borrada por la cuenta %s", member_id, getattr(actor, "pk", None))
    return member_id


def _ensure_photo_manager(actor, member):
    """La propia persona o quien gestiona el censo de su asociación."""
    if actor and actor.is_authenticated and member.account_id and member.account_id == actor.pk:
        return
    _ensure_manager(actor, member.association)


@transaction.atomic
def replace_member_photo(*, member, actor, upload):
    from .images import sanitize_member_photo

    locked = Member.objects.select_for_update().select_related("association").get(pk=member.pk)
    _ensure_photo_manager(actor, locked)
    sanitized = sanitize_member_photo(upload)
    old_name = locked.photo.name
    storage = locked.photo.storage
    locked.photo.save("photo.webp", sanitized, save=False)
    locked.updated_by = actor
    locked.save(update_fields=["photo", "updated_by", "updated_at"])
    if old_name:
        transaction.on_commit(lambda: storage.delete(old_name))
    return locked


@transaction.atomic
def delete_member_photo(*, member, actor):
    locked = Member.objects.select_for_update().select_related("association").get(pk=member.pk)
    _ensure_photo_manager(actor, locked)
    if locked.photo:
        locked.photo.delete(save=False)
        locked.updated_by = actor
        locked.save(update_fields=["photo", "updated_by", "updated_at"])
    return locked
